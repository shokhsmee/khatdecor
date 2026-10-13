"""One document of one event, and the Word file it turns into.

Filling a .docx means replacing ``{{marker}}`` text, and a marker survives in
one piece only while nobody has edited the paragraph: Word splits a paragraph
into runs on every formatting change, and a marker can end up cut across two
of them.  So the replacement works on the paragraph's whole text and writes
the result back into its first run — the paragraph keeps its own style, which
is where the formatting of these documents lives.
"""
import base64
import io
import logging

from odoo import _, api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

try:
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover - guarded by the manifest
    DocxDocument = None


class CstHrDocument(models.Model):
    _name = "cst.hr.document"
    _description = "Kadr hujjati"
    _order = "sequence, id"

    event_id = fields.Many2one(
        "cst.hr.event", string="Kadr harakati", required=True,
        ondelete="cascade", index=True)
    template_id = fields.Many2one(
        "cst.hr.doc.template", string="Shablon", required=True,
        ondelete="restrict")
    name = fields.Char(related="template_id.name", string="Hujjat", store=True)
    sequence = fields.Integer(related="template_id.sequence", store=True)
    employee_id = fields.Many2one(
        related="event_id.employee_id", string="Xodim", store=True, index=True)
    kind = fields.Selection(related="event_id.kind", string="Harakat turi", store=True)
    date = fields.Date(related="event_id.date", string="Sana", store=True)

    file = fields.Binary(string="Word fayli", attachment=True, copy=False)
    filename = fields.Char(string="Fayl nomi", copy=False)
    generated_on = fields.Datetime(string="Yaratilgan", readonly=True, copy=False)
    state = fields.Selection(
        [("todo", "Yaratilmagan"), ("ready", "Tayyor")],
        string="Holati", compute="_compute_state", store=True)

    @api.depends("file")
    def _compute_state(self):
        for document in self:
            document.state = "ready" if document.file else "todo"

    # ------------------------------------------------------------------
    def _placeholders(self):
        """What every ``{{marker}}`` in the template stands for."""
        self.ensure_one()
        event = self.event_id
        employee = event.employee_id.sudo()
        company = event.company_id
        date = event.date

        compensation = ""
        if event.kind == "boshash" and event.compensation_days > 0:
            compensation = _(
                "2. %(name)sga Mehnat kodeksining 234-moddasiga asosan "
                "foydalanilmagan barcha yillik asosiy ta'tillari uchun (%(days)s) "
                "kunlik kompensasiya to'lab berilsin",
                name=employee.name, days=event.compensation_days)

        return {
            "fish": employee.name or "",
            "hodisa": dict(event._fields["kind"].selection).get(event.kind, ""),
            "manzil": employee.cst_full_address or "",
            "passport": employee.passport_id or "",
            "jshshir": employee.cst_pinfl or "",
            "kim_bergan": employee.cst_passport_issuer or "",
            "telefon": employee.private_phone or employee.work_phone or "",
            "buyruq_no": event.order_no or "",
            "sana": date.strftime("%d.%m.%Y") if date else "",
            "sana_matn": event.date_text or "",
            "bolim": event.department_id.name or "",
            "lavozim": event.job_title or "",
            "shartnoma_no": event.contract_no or "",
            "yangi_bolim": event.new_department_id.name or "",
            "yangi_lavozim": event.new_job_title or "",
            "msh_sana": (event.base_contract_date.strftime("%d.%m.%Y")
                         if event.base_contract_date else ""),
            "msh_no": event.base_contract_no or "",
            "tatil_kuni": str(event.compensation_days or ""),
            "tatil_kompensatsiya": compensation,
            "kompaniya": company.cst_hr_company_label or company.name or "",
            "direktor": company.cst_hr_director or "",
            "shahar": company.cst_hr_place or "",
        }

    @staticmethod
    def _fill_paragraph(paragraph, values):
        text = paragraph.text
        if "{{" not in text:
            return
        filled = text
        for key, value in values.items():
            filled = filled.replace("{{%s}}" % key, value)
        if filled == text:
            return
        runs = paragraph.runs
        if not runs:
            paragraph.add_run(filled)
            return
        runs[0].text = filled
        for run in runs[1:]:
            run.text = ""

    @api.model
    def _fill_document(self, docx, values):
        for paragraph in docx.paragraphs:
            self._fill_paragraph(paragraph, values)
        for table in docx.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._fill_paragraph(paragraph, values)
        for section in docx.sections:
            for part in (section.header, section.footer):
                for paragraph in part.paragraphs:
                    self._fill_paragraph(paragraph, values)
        return docx

    def action_generate(self):
        if DocxDocument is None:
            raise UserError(_(
                "python-docx kutubxonasi o'rnatilmagan — Word hujjatlarini "
                "yaratib bo'lmaydi.  Server administratoriga murojaat qiling."))
        for document in self:
            template = document.template_id
            if not template.file:
                raise UserError(_(
                    "«%s» shabloniga Word fayli biriktirilmagan.", template.name))
            source = io.BytesIO(base64.b64decode(template.file))
            docx = document._fill_document(
                DocxDocument(source), document._placeholders())
            target = io.BytesIO()
            docx.save(target)
            employee = document.employee_id.sudo()
            slug = (employee.name or "xodim").replace(" ", "_").replace("/", "-")
            document.write({
                "file": base64.b64encode(target.getvalue()),
                "filename": f"{template.code}_{slug}.docx",
                "generated_on": fields.Datetime.now(),
            })
            _logger.info("HR document %s generated for %s",
                         template.code, employee.name)
        return True

    def action_download(self):
        self.ensure_one()
        if not self.file:
            self.action_generate()
        return {
            "type": "ir.actions.act_url",
            "url": (f"/web/content/cst.hr.document/{self.id}/file"
                    f"/{self.filename}?download=true"),
            "target": "self",
        }
