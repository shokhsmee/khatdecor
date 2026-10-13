"""Build the Word templates from the HR department's Excel workbook.

The wording of every document is the department's own, so it is not retyped
here: this reads «Xodimni ishga olish xujjatlari» and turns each document
sheet into a .docx whose formula cells have become ``{{placeholder}}`` markers.
Run it again after the department changes the workbook:

    python3 _build_templates.py /opt/kerama/docs/<workbook>.xlsx .
"""
import os
import re
import sys

import openpyxl
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

# Киритиш!<cell> → the placeholder the Odoo side fills in
FIELDS = {
    'B3': 'fish', 'C3': 'hodisa', 'H3': 'manzil', 'I3': 'passport', 'J3': 'jshshir',
    'K3': 'kim_bergan', 'L3': 'telefon', 'M3': 'buyruq_no', 'Q3': 'sana',
    'R3': 'sana_matn', 'S3': 'bolim', 'T3': 'lavozim', 'U3': 'shartnoma_no',
    'V3': 'yangi_bolim', 'W3': 'yangi_lavozim', 'X3': 'msh_sana', 'Y3': 'msh_no',
    'Z3': 'tatil_kuni',
}
# constants the sheets spell out that belong to the company record instead
COMPANY = {
    '"KHATDECOR" MCHJ direktori': '{{kompaniya}} direktori',
    'A.X.Tolipovga': '{{direktor}}ga',
    'A.X.Tolipov': '{{direktor}}',
    "Uchko'prik tumani": '{{shahar}}',
    '"KHATDECOR" MCHJ direktori:': '{{kompaniya}} direktori:',
    '"Khat decor" MCHJ direktori:': '{{kompaniya}} direktori:',
}
SHEETS = [
    ('Ариза', 'ariza_qabul', 'Ariza (ishga qabul)'),
    ('Buyruq', 'buyruq_qabul', "Buyruq (ishga qabul qilish to'g'risida)"),
    ('шарт', 'mehnat_shartnomasi', 'Mehnat shartnomasi'),
    ('RX', 'rozilik_xati', 'Rozilik xati'),
    ('QK', 'qoshimcha_kelishuv', "Qo'shimcha kelishuv"),
    ('Buyruq BL', 'buyruq_otkazish', "Buyruq (boshqa lavozimga o'tkazish)"),
    ("ariza bo'shash", 'ariza_boshash', 'Ariza (ishdan bo\'shash)'),
    ("Buyruq Bo'shash", 'buyruq_boshash', 'Buyruq (mehnat shartnomasini bekor qilish)'),
]


def resolve(ws, formula, seen=None):
    """Turn an Excel formula into text with ``{{placeholder}}`` markers."""
    seen = seen or set()
    f = formula.lstrip('=+').strip()

    m = re.fullmatch(r"IF\(Киритиш!Z3>0,\s*\+?CONCATENATE\(.*\),\s*\"\s*\"\s*,?\)", f, re.S)
    if m:
        return '{{tatil_kompensatsiya}}'

    if f.upper().startswith('CONCATENATE('):
        inner = f[f.index('(') + 1:f.rindex(')')]
        return ''.join(resolve(ws, part, seen) for part in split_args(inner))

    m = re.fullmatch(r"Киритиш!(\$?[A-Z]+\$?\d+)", f)
    if m:
        ref = m.group(1).replace('$', '')
        return '{{%s}}' % FIELDS.get(ref, 'noma_lum_%s' % ref)

    m = re.fullmatch(r"\$?([A-Z]+)\$?(\d+)", f)          # a self-reference
    if m:
        ref = f.replace('$', '')
        if ref in seen:
            return ''
        seen.add(ref)
        v = ws[ref].value
        return resolve(ws, v, seen) if isinstance(v, str) and v.startswith('=') else str(v or '')

    if f.startswith('"') and f.endswith('"'):
        return f[1:-1].replace('""', '"')
    return f


def split_args(s):
    """Split CONCATENATE arguments on top-level commas."""
    out, depth, quoted, cur = [], 0, False, ''
    for ch in s:
        if ch == '"':
            quoted = not quoted
        if not quoted:
            if ch == '(':
                depth += 1
            elif ch == ')':
                depth -= 1
            elif ch == ',' and depth == 0:
                out.append(cur.strip())
                cur = ''
                continue
        cur += ch
    if cur.strip():
        out.append(cur.strip())
    return out


def cell_text(ws, cell):
    v = cell.value
    if v is None:
        return ''
    if isinstance(v, str) and v.startswith('='):
        return resolve(ws, v)
    text = str(v)
    return COMPANY.get(text.strip(), text)


# Row blocks that are a real table on the paper sheet, with the Excel column
# ranges that make up each of their columns.
GRIDS = {
    'Buyruq': [(14, 24, [(1, 3), (4, 5), (6, 10)], True)],
    'Buyruq BL': [(13, 14, [(1, 3), (4, 4), (5, 8), (9, 10)], True)],
    "Buyruq Bo'shash": [(13, 14, [(1, 1), (2, 3), (4, 4), (5, 8), (9, 13)], True)],
    'шарт': [(80, 92, [(1, 6), (7, 11)], False)],
    'QK': [(22, 26, [(1, 6), (7, 11)], False)],
}
# rows whose text is the document's own heading
HEADINGS = {'Ариза': [9], 'Buyruq': [5, 9], 'шарт': [1, 9, 15, 19, 44, 51, 55, 58, 61, 68, 74, 79],
            'RX': [7], 'QK': [2, 11, 21], 'Buyruq BL': [5, 9],
            "ariza bo'shash": [7], "Buyruq Bo'shash": [5, 9]}
RIGHT = {'Ариза': range(1, 8), 'RX': range(1, 6), "ariza bo'shash": range(1, 6)}


def join_row(ws, row, lo=1, hi=None):
    hi = hi or ws.max_column
    parts = [cell_text(ws, ws.cell(row, c)) for c in range(lo, hi + 1)]
    return ' '.join(p for p in (p.strip() for p in parts) if p)


def add_par(doc, text, *, bold=False, align=None, size=10, space=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'
    return p


def build(ws, title, path):
    doc = Document()
    section = doc.sections[0]
    section.left_margin, section.right_margin = Cm(2.2), Cm(1.5)
    section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.5)
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(10)

    grids = GRIDS.get(ws.title, [])
    skip = {r for lo, hi, _, _ in grids for r in range(lo, hi + 1)}
    headings = set(HEADINGS.get(ws.title, []))
    right = set(RIGHT.get(ws.title, []))

    row = 1
    while row <= ws.max_row:
        grid = next((g for g in grids if g[0] == row), None)
        if grid:
            lo, hi, cols, borders = grid
            data = [[join_row(ws, r, a, b) for a, b in cols] for r in range(lo, hi + 1)]
            data = [r for r in data if any(r)]
            if data:
                table = doc.add_table(rows=len(data), cols=len(cols))
                table.style = 'Table Grid' if borders else 'Normal Table'
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for i, line in enumerate(data):
                    for j, text in enumerate(line):
                        cell = table.cell(i, j)
                        cell.text = ''
                        par = cell.paragraphs[0]
                        par.paragraph_format.space_after = Pt(1)
                        run = par.add_run(text)
                        run.font.size = Pt(9)
                        run.font.name = 'Times New Roman'
                doc.add_paragraph()
            row = hi + 1
            continue
        if row in skip:
            row += 1
            continue

        text = join_row(ws, row)
        if not text:
            add_par(doc, '', space=4)
            row += 1
            continue
        if row in headings:
            add_par(doc, text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11, space=6)
        elif row in right:
            add_par(doc, text, align=WD_ALIGN_PARAGRAPH.RIGHT)
        else:
            add_par(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        row += 1

    doc.save(path)
    return path


def main(src, outdir):
    wb = openpyxl.load_workbook(src, data_only=False)
    for sheet, code, title in SHEETS:
        path = os.path.join(outdir, code + '.docx')
        build(wb[sheet], title, path)
        print('%-24s -> %s' % (sheet, os.path.basename(path)))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else '.')
